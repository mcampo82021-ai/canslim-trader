import os
import openpyxl
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.discovery import build

SCOPES = ['https://www.googleapis.com/auth/spreadsheets', 'https://www.googleapis.com/auth/drive']
SPREADSHEET_ID = "1xxg81AY0DULhUcyRYZx2-gk5oByXe_ITXscMKR8w4Ns"
CREDS_FILE = os.path.expanduser('~/.config/google/canslim_credentials.json')
TOKEN_FILE = os.path.expanduser('~/.config/google/canslim_token.json')
EXCEL_FILE = os.path.expanduser('~/Documents/CAN SLIM/CAN_SLIM_Tracker_updated.xlsx')

def get_credentials():
    creds = None
    if os.path.exists(TOKEN_FILE):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(CREDS_FILE, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(TOKEN_FILE, 'w') as token:
            token.write(creds.to_json())
    return creds

def sync():
    print('📡 Leyendo Excel...')
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb[wb.sheetnames[0]]
    
    data = []
    for row in ws.iter_rows(values_only=True):
        data.append([str(cell) if cell is not None else '' for cell in row])
    
    print('🔑 Autenticando con Google...')
    creds = get_credentials()
    service = build('sheets', 'v4', credentials=creds)
    
    print('📤 Subiendo datos al Sheet...')
    sheet_name = ws.title
    range_name = f"'{sheet_name}'!A1"
    
    service.spreadsheets().values().clear(
        spreadsheetId=SPREADSHEET_ID,
        range=f"'{sheet_name}'!A1:AZ1000"
    ).execute()
    
    service.spreadsheets().values().update(
        spreadsheetId=SPREADSHEET_ID,
        range=range_name,
        valueInputOption='RAW',
        body={'values': data}
    ).execute()
    
    print(f'✅ Sincronizado — {len(data)} filas subidas al Google Sheet')

if __name__ == '__main__':
    sync()
